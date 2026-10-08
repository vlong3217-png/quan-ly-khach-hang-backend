from typing import Any, Dict, List, Optional
from fastapi import HTTPException, status
from app.core.security import hash_password
from app.schemas.user import UserCreate, UserStatusUpdate, UserUpdate
from app.services.auth_service import fake_users_db
from app.services.data_handover_service import execute_handover

from app.core.database import SessionLocal
from app.models.user import User as UserModel

ALLOWED_ROLES = {"ADMIN", "MANAGER", "USER"}
ALLOWED_STATUSES = {"ACTIVE", "LOCKED"}


def get_all_users(
    search: Optional[str] = None,
    role: Optional[str] = None,
    is_active: Optional[bool] = None,
    skip: int = 0,
    limit: int = 20,
) -> List[dict]:
    results = fake_users_db

    if search and search.strip():
        term = search.strip().lower()
        results = [
            u for u in results
            if term in u["email"].lower()
            or term in u.get("full_name", "").lower()
            or (u.get("username") and term in u["username"].lower())
        ]

    if role and role.strip():
        role_clean = role.strip().upper()
        results = [u for u in results if u.get("role", "").upper() == role_clean]

    if is_active is not None:
        results = [u for u in results if u.get("is_active") == is_active]

    return results[skip : skip + limit]


def get_user_by_id(user_id: int) -> dict:
    for u in fake_users_db:
        if u["id"] == user_id:
            if "status" not in u:
                u["status"] = "ACTIVE" if u.get("is_active", True) else "LOCKED"
            return u
    raise HTTPException(
        status_code=status.HTTP_404_NOT_FOUND,
        detail=f"Không tìm thấy tài khoản người dùng với id {user_id}",
    )


def create_user(user_in: UserCreate) -> dict:
    clean_email = user_in.email.strip().lower()

    # Check email duplicate
    for u in fake_users_db:
        if u["email"].lower() == clean_email:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Email đã được sử dụng",
            )

    # Check username duplicate if provided
    clean_username = user_in.username.strip().lower() if user_in.username else None
    if clean_username:
        for u in fake_users_db:
            if u.get("username") and u["username"].lower() == clean_username:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Tên đăng nhập đã được sử dụng",
                )

    # Validate role
    role = (user_in.role or "USER").strip().upper()
    if role not in ALLOWED_ROLES:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Vai trò không hợp lệ. Các vai trò cho phép: {', '.join(sorted(ALLOWED_ROLES))}",
        )

    hashed = hash_password(user_in.password)

    new_id = max([u["id"] for u in fake_users_db], default=0) + 1
    is_active = True if user_in.is_active is None else user_in.is_active
    new_user = {
        "id": new_id,
        "email": user_in.email.strip(),
        "username": user_in.username.strip() if user_in.username else clean_email.split("@")[0],
        "full_name": user_in.full_name.strip(),
        "role": role,
        "is_active": is_active,
        "status": "ACTIVE" if is_active else "LOCKED",
        "hashed_password": hashed,
    }

    fake_users_db.append(new_user)

    # Đồng bộ lưu vào CSDL MySQL
    try:
        db = SessionLocal()
        db_user = UserModel(
            id=new_id,
            email=new_user["email"],
            username=new_user["username"],
            full_name=new_user["full_name"],
            role=new_user["role"],
            hashed_password=new_user["hashed_password"],
            is_active=new_user["is_active"],
            team_id=new_user.get("team_id"),
            status=new_user["status"],
        )
        db.merge(db_user)
        db.commit()
        db.close()
    except Exception:
        pass

    return new_user


def update_user(user_id: int, user_in: UserUpdate) -> dict:
    user = get_user_by_id(user_id)

    # Email update & duplicate check
    if user_in.email is not None:
        clean_email = user_in.email.strip().lower()
        for u in fake_users_db:
            if u["id"] != user_id and u["email"].lower() == clean_email:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Email đã được sử dụng bởi tài khoản khác",
                )
        user["email"] = user_in.email.strip()

    # Username update & duplicate check
    if user_in.username is not None:
        clean_username = user_in.username.strip().lower()
        for u in fake_users_db:
            if u["id"] != user_id and u.get("username") and u["username"].lower() == clean_username:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Tên đăng nhập đã được sử dụng bởi tài khoản khác",
                )
        user["username"] = user_in.username.strip()

    # Full name update
    if user_in.full_name is not None:
        if not user_in.full_name.strip():
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Họ và tên không được để trống",
            )
        user["full_name"] = user_in.full_name.strip()

    # Role update
    if user_in.role is not None:
        role = user_in.role.strip().upper()
        if role not in ALLOWED_ROLES:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Vai trò không hợp lệ. Các vai trò cho phép: {', '.join(sorted(ALLOWED_ROLES))}",
            )
        user["role"] = role

    # is_active update
    if user_in.is_active is not None:
        user["is_active"] = user_in.is_active
        user["status"] = "ACTIVE" if user_in.is_active else "LOCKED"

    # Password update
    if user_in.password is not None:
        if len(user_in.password.strip()) < 6:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Mật khẩu phải có ít nhất 6 ký tự",
            )
        user["hashed_password"] = hash_password(user_in.password.strip())

    # Đồng bộ cập nhật vào CSDL MySQL
    try:
        db = SessionLocal()
        db_user = db.query(UserModel).filter(UserModel.id == user_id).first()
        if db_user:
            if user_in.email is not None:
                db_user.email = user["email"]
            if user_in.username is not None:
                db_user.username = user["username"]
            if user_in.full_name is not None:
                db_user.full_name = user["full_name"]
            if user_in.role is not None:
                db_user.role = user["role"]
            if user_in.is_active is not None:
                db_user.is_active = user["is_active"]
                db_user.status = user["status"]
            if user_in.password is not None:
                db_user.hashed_password = user["hashed_password"]
            db.commit()
        db.close()
    except Exception:
        pass

    return user


def update_user_status(
    user_id: int,
    status_in: UserStatusUpdate,
    current_admin: dict,
) -> Dict[str, Any]:
    """
    Cập nhật trạng thái tài khoản: ACTIVE hoặc LOCKED.
    - Không cho phép Admin tự khóa tài khoản của chính mình.
    - Khi khóa (LOCKED): Đặt is_active = False, status = LOCKED.
    - Hỗ trợ bàn giao dữ liệu nếu có chỉ định handover_to_user_id.
    - Khi mở khóa (ACTIVE): Đặt is_active = True, status = ACTIVE.
    """
    user = get_user_by_id(user_id)

    status_upper = status_in.status.strip().upper()
    if status_upper not in ALLOWED_STATUSES:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Trạng thái không hợp lệ. Chỉ chấp nhận các giá trị: {', '.join(sorted(ALLOWED_STATUSES))}",
        )

    if status_upper == "LOCKED" and current_admin.get("id") == user_id:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Không thể tự khóa tài khoản của chính mình",
        )

    handover_result = None

    if status_upper == "LOCKED":
        user["status"] = "LOCKED"
        user["is_active"] = False
        message = f"Đã khóa tài khoản thành công cho user #{user_id}"

        if status_in.handover_to_user_id is not None:
            handover_result = execute_handover(
                source_user_id=user_id,
                target_user_id=status_in.handover_to_user_id,
            )
            message += f" và bàn giao dữ liệu sang user #{status_in.handover_to_user_id}"
    else:
        user["status"] = "ACTIVE"
        user["is_active"] = True
        message = f"Đã mở khóa tài khoản thành công cho user #{user_id}"

    return {
        "id": user["id"],
        "email": user["email"],
        "username": user.get("username"),
        "full_name": user["full_name"],
        "role": user["role"],
        "is_active": user["is_active"],
        "status": user["status"],
        "message": message,
        "handover": handover_result,
    }


def generate_user_template_excel() -> bytes:
    """Tạo file Excel mẫu để nhập người dùng hàng loạt."""
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Danh sách người dùng"

    headers = ["Email (*)", "Họ và tên (*)", "Mật khẩu (*)", "Vai trò", "Team ID"]
    ws.append(headers)

    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    # Add sample rows
    samples = [
        ["nguyen.van.a@example.com", "Nguyễn Văn A", "Pass@1234", "USER", 1],
        ["tran.thi.b@example.com", "Trần Thị B", "Pass@1234", "MANAGER", 2],
        ["le.van.c@example.com", "Lê Văn C", "Pass@1234", "ADMIN", ""],
    ]
    for row in samples:
        ws.append(row)

    for row in ws.iter_rows(min_row=2, max_row=len(samples) + 1, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    col_widths = [32, 28, 20, 16, 14]
    for col_idx, width in enumerate(col_widths, start=1):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def parse_and_validate_user_rows(file_contents: bytes) -> Dict[str, Any]:
    """
    Đọc tệp Excel và kiểm tra tính hợp lệ của từng dòng:
    - Báo lỗi chi tiết cho từng dòng nếu thiếu thông tin, sai định dạng email, mật khẩu ngắn, vai trò không tồn tại, hoặc trùng lặp email.
    """
    import io
    import re
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_contents), data_only=True)
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Tệp không đúng định dạng Excel hợp lệ: {str(e)}",
        )

    ws = wb.active
    rows_data = []
    row_idx = 0

    seen_emails = set()
    existing_db_emails = {u["email"].lower() for u in fake_users_db}

    email_regex = re.compile(r"^[\w\.-]+@[\w\.-]+\.\w+$")

    for row in ws.iter_rows(values_only=True):
        row_idx += 1
        if row_idx == 1:
            # Skip header row
            continue

        if not any(row):
            # Skip entirely empty rows
            continue

        email_val = str(row[0]).strip() if len(row) > 0 and row[0] is not None else ""
        full_name_val = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
        password_val = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
        role_val = str(row[3]).strip().upper() if len(row) > 3 and row[3] is not None else "USER"
        team_id_raw = row[4] if len(row) > 4 else None

        team_id_val = None
        if team_id_raw is not None and str(team_id_raw).strip() != "":
            try:
                team_id_val = int(team_id_raw)
            except ValueError:
                team_id_val = "INVALID"

        errors = []

        # 1. Email check
        if not email_val:
            errors.append("Email không được để trống")
        elif not email_regex.match(email_val):
            errors.append("Email không đúng định dạng")
        elif email_val.lower() in existing_db_emails:
            errors.append(f"Email '{email_val}' đã tồn tại trong hệ thống")
        elif email_val.lower() in seen_emails:
            errors.append(f"Email '{email_val}' bị trùng lặp trong tệp tải lên")
        else:
            seen_emails.add(email_val.lower())

        # 2. Full name check
        if not full_name_val:
            errors.append("Họ và tên không được để trống")

        # 3. Password check
        if not password_val:
            errors.append("Mật khẩu không được để trống")
        elif len(password_val) < 6:
            errors.append("Mật khẩu phải có tối thiểu 6 ký tự")

        # 4. Role check
        if not role_val:
            role_val = "USER"
        elif role_val not in ALLOWED_ROLES:
            errors.append(f"Vai trò '{role_val}' không hợp lệ. Các vai trò hợp lệ: {', '.join(sorted(ALLOWED_ROLES))}")

        # 5. Team ID check
        if team_id_val == "INVALID":
            errors.append("Team ID phải là số nguyên")
        elif team_id_val is not None and (team_id_val < 1 or team_id_val > 10):
            errors.append("Team ID không hợp lệ trong hệ thống")

        # AC S1-09 rule: Nếu role là MANAGER thì bắt buộc có team_id
        if role_val == "MANAGER" and not team_id_val:
            errors.append("Vai trò MANAGER bắt buộc phải thuộc về một nhóm kinh doanh (Team ID)")

        is_valid = len(errors) == 0

        rows_data.append({
            "row_number": row_idx,
            "email": email_val or None,
            "full_name": full_name_val or None,
            "password": password_val or None,
            "role": role_val,
            "team_id": team_id_val if isinstance(team_id_val, int) else None,
            "is_valid": is_valid,
            "errors": errors,
        })

    valid_count = sum(1 for r in rows_data if r["is_valid"])
    invalid_count = len(rows_data) - valid_count

    return {
        "total_rows": len(rows_data),
        "valid_rows_count": valid_count,
        "invalid_rows_count": invalid_count,
        "rows": rows_data,
    }


def execute_user_import(file_contents: bytes) -> Dict[str, Any]:
    """
    Nhập người dùng hàng loạt từ Excel:
    - Bỏ qua các dòng bị lỗi.
    - Nhập thành công các dòng hợp lệ.
    - Trả về báo cáo tổng kết chi tiết.
    """
    analysis = parse_and_validate_user_rows(file_contents)
    details = []
    imported_count = 0
    skipped_count = 0

    max_id = max([u["id"] for u in fake_users_db], default=0)

    for item in analysis["rows"]:
        row_num = item["row_number"]
        if not item["is_valid"]:
            skipped_count += 1
            details.append({
                "row_number": row_num,
                "email": item["email"],
                "full_name": item["full_name"],
                "status": "SKIPPED",
                "error_message": "; ".join(item["errors"]),
            })
            continue

        # Valid row -> Create user
        max_id += 1
        new_user = {
            "id": max_id,
            "email": item["email"].lower(),
            "username": item["email"].split("@")[0],
            "full_name": item["full_name"],
            "hashed_password": hash_password(item["password"]),
            "role": item["role"],
            "team_id": item["team_id"],
            "is_active": True,
            "status": "ACTIVE",
            "token_version": 0,
            "failed_login_attempts": 0,
            "temporary_lock_until": None,
        }
        fake_users_db.append(new_user)
        try:
            db = SessionLocal()
            db_u = UserModel(
                id=max_id,
                email=new_user["email"],
                username=new_user["username"],
                full_name=new_user["full_name"],
                role=new_user["role"],
                hashed_password=new_user["hashed_password"],
                is_active=new_user["is_active"],
                team_id=new_user["team_id"],
                status=new_user["status"],
            )
            db.merge(db_u)
            db.commit()
            db.close()
        except Exception:
            pass

        imported_count += 1
        details.append({
            "row_number": row_num,
            "email": item["email"],
            "full_name": item["full_name"],
            "status": "SUCCESS",
            "error_message": None,
        })

    return {
        "total_rows": analysis["total_rows"],
        "imported_count": imported_count,
        "skipped_count": skipped_count,
        "details": details,
    }


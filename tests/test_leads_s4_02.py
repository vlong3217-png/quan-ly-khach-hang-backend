import io
import pytest
import openpyxl
from fastapi.testclient import TestClient

from app.main import app
from app.services.auth_service import reset_fake_users
from app.services.customer_service import reset_fake_customers
from app.services import lead_service

client = TestClient(app)


def get_auth_token(email: str = "admin@gmail.com", password: str = "123456") -> str:
    response = client.post(
        "/auth/login",
        json={"email": email, "password": password},
    )
    assert response.status_code == 200, f"Login failed: {response.text}"
    return response.json()["access_token"]


@pytest.fixture(autouse=True)
def setup_data():
    reset_fake_users()
    reset_fake_customers()
    lead_service.reset_fake_leads()
    yield
    reset_fake_users()
    reset_fake_customers()
    lead_service.reset_fake_leads()


# ==============================================================================
# 1. AC S4-02: Nhập tay một lead từ sự kiện hoặc danh thiếp & Bắt buộc có nguồn
# ==============================================================================

def test_manual_create_lead_success_with_mandatory_source():
    """Nhập tay một lead từ sự kiện hoặc danh thiếp thành công khi có đủ nguồn."""
    token = get_auth_token()
    headers = {"Authorization": f"Bearer {token}"}

    payload = {
        "name": "Đặng Hoàng Giang",
        "company_name": "Công ty TNHH Giải pháp Đám mây Giang Nam",
        "title": "Giám đốc Công nghệ (CTO)",
        "email": "giang.dang@giangnam.vn",
        "phone": "0988776655",
        "address": "Tầng 12, Keangnam Landmark 72, Hà Nội",
        "source": "Hội thảo Tech Expo 2026",
        "notes": "Nhận danh thiếp trực tiếp tại quầy triển lãm",
    }
    resp = client.post("/leads", json=payload, headers=headers)
    assert resp.status_code == 201
    data = resp.json()
    assert data["name"] == "Đặng Hoàng Giang"
    assert data["source"] == "Hội thảo Tech Expo 2026"
    assert data["company_name"] == "Công ty TNHH Giải pháp Đám mây Giang Nam"
    assert data["id"] is not None
    assert data["status"] == "NEW"


def test_manual_create_lead_missing_source_rejected_422():
    """Mọi lead nhập vào đều bắt buộc có nguồn. Thiếu nguồn bị từ chối 422."""
    token = get_auth_token()
    headers = {"Authorization": f"Bearer {token}"}

    # Trường hợp 1: Không có key 'source'
    payload_no_source = {
        "name": "Nguyễn Văn Thiếu Nguồn",
        "phone": "0912345678",
    }
    resp1 = client.post("/leads", json=payload_no_source, headers=headers)
    assert resp1.status_code == 422

    # Trường hợp 2: Có key 'source' nhưng giá trị rỗng hoặc toàn khoảng trắng
    payload_blank_source = {
        "name": "Nguyễn Văn Nguồn Rỗng",
        "phone": "0912345678",
        "source": "   ",
    }
    resp2 = client.post("/leads", json=payload_blank_source, headers=headers)
    assert resp2.status_code == 422


def test_manual_create_lead_missing_name_rejected_422():
    """Họ và tên lead là bắt buộc."""
    token = get_auth_token()
    headers = {"Authorization": f"Bearer {token}"}

    payload = {
        "source": "Sự kiện",
        "phone": "0912345678",
    }
    resp = client.post("/leads", json=payload, headers=headers)
    assert resp.status_code == 422


# ==============================================================================
# 2. AC S4-02: Nhập hàng loạt có tệp mẫu Excel
# ==============================================================================

def test_download_lead_import_template():
    """Tải tệp mẫu Excel có các cột rõ ràng và hàng minh họa."""
    token = get_auth_token()
    headers = {"Authorization": f"Bearer {token}"}

    resp = client.get("/leads/import/template", headers=headers)
    assert resp.status_code == 200
    assert "application/vnd.openxmlformats" in resp.headers["content-type"]
    assert len(resp.content) > 0

    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert ws.title in ["Mau_Nhap_Lead", "Lead_Template"]
    assert ws.cell(row=1, column=1).value == "Họ và tên (*)"
    assert ws.cell(row=1, column=2).value == "Nguồn lead (*)"
    assert ws.cell(row=1, column=3).value == "Số điện thoại"
    assert ws.cell(row=1, column=4).value == "Email"
    assert "Huong_Dan" in wb.sheetnames


def test_download_lead_import_template_is_sample_not_from_system():
    """Tải tệp mẫu không cần token (trực tiếp trình duyệt) và xác nhận là mẫu minh họa, không lấy từ DB."""
    resp = client.get("/leads/import/template")
    assert resp.status_code == 200
    assert "application/vnd.openxmlformats" in resp.headers["content-type"]
    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb["Mau_Nhap_Lead"]
    # Xác nhận chỉ có header và đúng 2 dòng ví dụ minh họa
    assert ws.max_row == 3  # row 1 header + 2 rows sample
    assert ws.cell(row=2, column=1).value == "Trần Quốc Bảo"
    assert ws.cell(row=3, column=1).value == "Lê Thu Hà"
    # Xác nhận có sheet hướng dẫn
    assert "Huong_Dan" in wb.sheetnames


# ==============================================================================
# 3. AC S4-02: Xem trước và báo lỗi theo từng dòng (Row-level validation)
# ==============================================================================

def test_preview_lead_excel_import_with_row_level_errors_and_duplicates():
    """Xem trước file Excel, phát hiện lỗi cụ thể theo từng dòng và phát hiện trùng lặp."""
    token = get_auth_token()
    headers = {"Authorization": f"Bearer {token}"}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Họ và tên (*)", "Nguồn lead (*)", "Số điện thoại", "Email", "Tên công ty", "Chức vụ", "Địa chỉ", "Ghi chú", "Mã chiến dịch"])
    
    # Dòng 2: Hợp lệ 100%
    ws.append(["Bùi Minh Trí", "Hội thảo Fintech 2026", "0918889999", "tri.bui@fintech.vn", "Công ty Fintech Á Châu", "Trưởng dự án", "Hà Nội", "Trao đổi danh thiếp", ""])
    
    # Dòng 3: Thiếu nguồn lead (Lỗi AC S4-02)
    ws.append(["Phạm Thị Lan", "", "0917778888", "lan.pham@gmail.com", "Lan Boutique", "Chủ cửa hàng", "", "", ""])
    
    # Dòng 4: Thiếu họ tên lead (Lỗi)
    ws.append(["", "Sự kiện Kết nối", "0916667777", "contact@startup.vn", "Startup Việt", "Founder", "", "", ""])
    
    # Dòng 5: Sai định dạng email (Lỗi)
    ws.append(["Hoàng Văn Đức", "Danh thiếp", "0915556666", "email-sai-dinh-dang", "Công ty Cơ khí Đức", "Giám đốc", "", "", ""])
    
    # Dòng 6: Trùng số điện thoại với Lead 1 đã có ("0912345678" của Trần Văn Hùng)
    ws.append(["Trần Văn Hùng Cập Nhật", "Hội nghị", "0912345678", "hung.tran@smarttech.vn", "SmartTech", "Trưởng phòng", "", "", ""])

    buffer = io.BytesIO()
    wb.save(buffer)
    excel_bytes = buffer.getvalue()

    files = {"file": ("test_leads_import.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    resp = client.post("/leads/import/preview", files=files, headers=headers)
    assert resp.status_code == 200
    pdata = resp.json()

    assert pdata["total_rows"] == 5
    assert pdata["valid_rows_count"] == 2  # Dòng 2 và Dòng 6 (Dòng 6 hợp lệ cú pháp nhưng là trùng)
    assert pdata["invalid_rows_count"] == 3
    assert pdata["duplicate_rows_count"] >= 1

    # Kiểm tra lỗi từng dòng
    row2 = next(r for r in pdata["rows"] if r["row_number"] == 2)
    assert row2["is_valid"] is True
    assert len(row2["errors"]) == 0

    row3 = next(r for r in pdata["rows"] if r["row_number"] == 3)
    assert row3["is_valid"] is False
    assert any("Nguồn lead là bắt buộc" in err for err in row3["errors"])

    row4 = next(r for r in pdata["rows"] if r["row_number"] == 4)
    assert row4["is_valid"] is False
    assert any("Họ và tên" in err for err in row4["errors"])

    row5 = next(r for r in pdata["rows"] if r["row_number"] == 5)
    assert row5["is_valid"] is False
    assert any("Email không đúng định dạng" in err for err in row5["errors"])

    row6 = next(r for r in pdata["rows"] if r["row_number"] == 6)
    assert row6["is_duplicate"] is True
    assert row6["existing_lead_id"] == 1


# ==============================================================================
# 4. AC S4-02: Nhập dữ liệu sau khi xem trước (Commit Import)
# ==============================================================================

def test_commit_lead_excel_import_with_skip_duplicate():
    """Commit nhập file Excel với chế độ SKIP dòng trùng lặp."""
    token = get_auth_token()
    headers = {"Authorization": f"Bearer {token}"}

    rows = [
        {
            "row_number": 2,
            "name": "Vũ Đình Quân",
            "source": "Hội thảo Bảo mật",
            "phone": "0934567890",
            "email": "quan.vu@security.vn",
            "company_name": "Công ty CyberSafe",
            "is_valid": True,
            "is_duplicate": False,
            "errors": [],
            "duplicate_reasons": [],
        },
        {
            "row_number": 3,
            "name": "Trần Văn Hùng Trùng",
            "source": "Sự kiện",
            "phone": "0912345678",
            "email": "hung.tran@smarttech.vn",
            "is_valid": True,
            "is_duplicate": True,
            "existing_lead_id": 1,
            "errors": [],
            "duplicate_reasons": ["Trùng số điện thoại: 0912345678"],
        },
        {
            "row_number": 4,
            "name": "",
            "source": "",
            "is_valid": False,
            "is_duplicate": False,
            "errors": ["Họ và tên lead là bắt buộc"],
            "duplicate_reasons": [],
        },
    ]

    commit_payload = {
        "duplicate_handling": "SKIP",
        "rows": rows,
    }
    resp = client.post("/leads/import/commit", json=commit_payload, headers=headers)
    assert resp.status_code == 200
    data = resp.json()

    assert data["inserted_count"] == 1
    assert data["skipped_count"] == 1
    assert data["failed_count"] == 1

    # Lead mới được thêm vào danh sách
    leads = client.get("/leads", headers=headers).json()["leads"]
    assert any(l["name"] == "Vũ Đình Quân" for l in leads)


def test_commit_lead_excel_import_with_update_duplicate():
    """Commit nhập file Excel với chế độ UPDATE dòng trùng lặp."""
    token = get_auth_token()
    headers = {"Authorization": f"Bearer {token}"}

    rows = [
        {
            "row_number": 2,
            "name": "Trần Văn Hùng Đã Cập Nhật",
            "source": "Hội thảo Quốc tế",
            "phone": "0912345678",
            "email": "hung.tran@smarttech.vn",
            "company_name": "SmartTech Global",
            "title": "Tổng Giám Đốc",
            "is_valid": True,
            "is_duplicate": True,
            "existing_lead_id": 1,
            "errors": [],
            "duplicate_reasons": ["Trùng số điện thoại"],
        }
    ]

    commit_payload = {
        "duplicate_handling": "UPDATE",
        "rows": rows,
    }
    resp = client.post("/leads/import/commit", json=commit_payload, headers=headers)
    assert resp.status_code == 200
    data = resp.json()

    assert data["updated_count"] == 1

    # Kiểm tra lead ID 1 đã được cập nhật
    lead1 = client.get("/leads/1", headers=headers).json()
    assert lead1["name"] == "Trần Văn Hùng Đã Cập Nhật"
    assert lead1["company_name"] == "SmartTech Global"
    assert lead1["title"] == "Tổng Giám Đốc"

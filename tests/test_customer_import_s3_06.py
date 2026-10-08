import io
import pytest
import openpyxl
from fastapi.testclient import TestClient

from app.main import app
from app.services.auth_service import reset_fake_users
from app.services.contact_service import reset_fake_contacts
from app.services import customer_service

client = TestClient(app)



def get_auth_token(email: str = "admin@gmail.com", password: str = "123456") -> str:
    response = client.post(
        "/auth/login",
        json={"email": email, "password": password},
    )
    assert response.status_code == 200, f"Login failed: {response.text}"
    return response.json()["access_token"]


@pytest.fixture(autouse=True)
def setup_data():
    reset_fake_users()
    customer_service.reset_fake_customers()
    reset_fake_contacts()
    yield
    reset_fake_users()
    customer_service.reset_fake_customers()
    reset_fake_contacts()


def test_download_customer_import_template():
    token = get_auth_token("admin@gmail.com", "123456")
    headers = {"Authorization": f"Bearer {token}"}

    resp = client.get("/customers/import/template", headers=headers)
    assert resp.status_code == 200
    assert "application/vnd.openxmlformats" in resp.headers["content-type"]
    assert len(resp.content) > 0

    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    assert ws.title == "KhachHang_Template"
    assert ws.cell(row=1, column=1).value == "Tên khách hàng (*)"


def test_preview_and_commit_import_excel_with_skip_duplicate():
    token = get_auth_token("admin@gmail.com", "123456")
    headers = {"Authorization": f"Bearer {token}"}

    # Tạo file Excel test có:
    # 1 dòng mới hoàn toàn
    # 1 dòng trùng MST với Customer 1 ("0101234567")
    # 1 dòng không có tên khách hàng (lỗi)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Tên", "MST", "Ngành", "Quy mô", "Website", "Địa chỉ", "SĐT", "Email", "Status"])
    ws.append(["Doanh nghiệp mới XYZ", "0109990001", "Công nghệ", "50 nhân sự", "https://xyz.vn", "Hà Nội", "09111", "m@xyz.vn", "PROSPECT"])
    ws.append(["Doanh nghiệp trùng MST", "0101234567", "CNTT", "100 nhân sự", "https://abc.vn", "HN", "09222", "m@abc.vn", "CUSTOMER"])
    ws.append(["", "0108880002", "Dịch vụ", "20 nhân sự", "", "", "", "", "PROSPECT"])

    buffer = io.BytesIO()
    wb.save(buffer)
    excel_bytes = buffer.getvalue()

    # 1. Preview
    files = {"file": ("test_import.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    preview_resp = client.post("/customers/import/preview", files=files, headers=headers)
    assert preview_resp.status_code == 200
    pdata = preview_resp.json()
    assert pdata["total_rows"] == 3
    assert pdata["duplicate_rows_count"] == 1
    assert pdata["invalid_rows_count"] == 1

    # 2. Commit với tùy chọn SKIP
    commit_payload = {
        "duplicate_handling": "SKIP",
        "rows": pdata["rows"],
    }
    commit_resp = client.post("/customers/import/commit", json=commit_payload, headers=headers)
    assert commit_resp.status_code == 200
    cdata = commit_resp.json()
    assert cdata["inserted_count"] == 1
    assert cdata["skipped_count"] == 1
    assert cdata["failed_count"] == 1

    # Kiểm tra xem doanh nghiệp mới đã có trong DB chưa
    assert any(c["name"] == "Doanh nghiệp mới XYZ" for c in customer_service.FAKE_CUSTOMERS)


def test_commit_import_excel_with_update_duplicate():
    token = get_auth_token("admin@gmail.com", "123456")
    headers = {"Authorization": f"Bearer {token}"}

    # Tạo dòng trùng MST với Customer 1 ("0101234567") nhưng thay đổi địa chỉ
    row = {
        "row_number": 2,
        "name": "Công ty ABC Cập Nhật Địa Chỉ",
        "tax_code": "0101234567",
        "address": "Địa chỉ mới số 999 Phố Huế",
        "is_valid": True,
        "is_duplicate": True,
        "existing_customer_id": 1,
    }

    commit_payload = {
        "duplicate_handling": "UPDATE",
        "rows": [row],
    }
    commit_resp = client.post("/customers/import/commit", json=commit_payload, headers=headers)
    assert commit_resp.status_code == 200
    cdata = commit_resp.json()
    assert cdata["updated_count"] == 1

    # Kiểm tra xem Customer 1 đã cập nhật chưa
    cust1 = next(c for c in customer_service.FAKE_CUSTOMERS if c["id"] == 1)
    assert cust1["address"] == "Địa chỉ mới số 999 Phố Huế"
    assert cust1["name"] == "Công ty ABC Cập Nhật Địa Chỉ"

